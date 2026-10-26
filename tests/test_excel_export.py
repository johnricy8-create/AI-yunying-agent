from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from ecom_ops.exceptions import HistoryError
from ecom_ops.history.excel import export_history_excel
from ecom_ops.history.store import record_result


def _result() -> dict:
    metrics = pd.DataFrame(
        [
            {
                "product_id": "P1",
                "product_name": "测试商品A",
                "shop": "店铺A",
                "category": "Excel模板",
                "pv": 120.0,
                "uv": 100.0,
                "orders": 0.0,
                "sales_amount": 0.0,
                "cr_today": 0.0,
                "pv_change_rate": 0.1,
            }
        ]
    )
    return {
        "run_id": "excel-run",
        "stat_date": "2026-07-31",
        "source_path": "test.csv",
        "report_path": "reports/test.md",
        "status": "completed",
        "metrics": metrics,
        "anomalies": [
            {
                "product_id": "P1",
                "anomaly_types": ["零转化"],
                "severity_score": 100.0,
                "severity_level": "高危",
                "low_confidence": False,
                "pv_baseline": 100.0,
                "orders_baseline": 3.0,
            }
        ],
        "global_summary": {
            "product_count": 1,
            "total_pv": 120,
            "total_orders": 0,
            "total_sales": 0,
            "overall_cr": 0.0,
            "overall_pv_change_rate": 0.1,
        },
        "ai_result": {
            "status": "ok",
            "products": {
                "P1": {
                    "recommendations": [
                        {
                            "action": "检查详情页",
                            "rationale": "零转化",
                            "risk": "无",
                            "needs_human": True,
                        }
                    ]
                }
            },
        },
        "data_quality": {"summary": {"valid_rows": 1}},
        "data_note": "Excel 测试",
    }


def test_export_history_excel(settings, tmp_path):
    history_settings = replace(
        settings,
        history=replace(
            settings.history,
            enabled=True,
            db_path=tmp_path / "history.db",
        ),
    )
    record_result(history_settings, _result())
    output = export_history_excel(history_settings, tmp_path / "export.xlsx", days=7)

    workbook = load_workbook(output)
    assert {
        "说明",
        "每日概览",
        "商品指标",
        "异常明细",
        "建议清单",
        "建议追踪",
    }.issubset(set(workbook.sheetnames))
    assert workbook["商品指标"].max_row >= 2
    assert workbook["异常明细"].max_row >= 2
    assert workbook["建议清单"].max_row >= 2
    assert workbook["每日概览"]["A2"].value == "2026-07-31"


def test_export_missing_history_raises(settings, tmp_path):
    history_settings = replace(
        settings,
        history=replace(
            settings.history,
            enabled=True,
            db_path=tmp_path / "missing.db",
        ),
    )
    with pytest.raises(HistoryError):
        export_history_excel(history_settings, tmp_path / "export.xlsx")
