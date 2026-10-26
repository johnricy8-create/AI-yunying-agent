"""真实数据模板生成。

模板表头使用中文别名，加载器会通过 config/fields.yaml 映射为内部字段。
Excel 额外包含“字段说明”工作表。
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ecom_ops.config import FieldsConfig
from ecom_ops.utils.files import ensure_dir

TEMPLATE_HEADERS = [
    "统计日期",
    "商品ID",
    "商品名称",
    "商品简介",
    "店铺",
    "平台",
    "类目",
    "单价",
    "成本",
    "库存",
    "浏览量",
    "访客数",
    "订单量",
    "销售额",
    "退款金额",
]

NOTE_HEADERS = ["字段名", "中文名", "类型", "必填", "说明", "示例"]

EXAMPLE_VALUES = {
    "stat_date": "2026-10-25",
    "product_id": "P0001",
    "product_name": "销售报表模板",
    "product_desc": "可编辑 Excel 模板，含使用说明",
    "shop": "旗舰店",
    "platform": "淘宝",
    "category": "Excel模板",
    "price": "19.90",
    "cost": "5.00",
    "stock": "100",
    "pv": "300",
    "uv": "250",
    "orders": "8",
    "sales_amount": "159.20",
    "refund_amount": "0",
}


def create_csv_template(path: str | Path) -> Path:
    file_path = Path(path)
    ensure_dir(file_path.parent)
    pd.DataFrame(columns=TEMPLATE_HEADERS).to_csv(
        file_path,
        index=False,
        encoding="utf-8-sig",
    )
    return file_path


def _style_header(row) -> None:
    fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in row:
        cell.font = Font(bold=True)
        cell.fill = fill


def create_excel_template(path: str | Path, fields: FieldsConfig) -> Path:
    file_path = Path(path)
    ensure_dir(file_path.parent)

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "销售数据"
    data_sheet.append(TEMPLATE_HEADERS)
    _style_header(data_sheet[1])
    data_sheet.freeze_panes = "A2"
    for index, header in enumerate(TEMPLATE_HEADERS, start=1):
        data_sheet.column_dimensions[get_column_letter(index)].width = max(
            12, min(30, len(header) * 2 + 8)
        )

    note_sheet = workbook.create_sheet("字段说明")
    note_sheet.append(NOTE_HEADERS)
    _style_header(note_sheet[1])
    note_sheet.freeze_panes = "A2"
    for spec in fields.fields.values():
        note_sheet.append(
            [
                spec.name,
                spec.label,
                spec.dtype,
                "是" if spec.required else "否",
                spec.description,
                EXAMPLE_VALUES.get(spec.name, ""),
            ]
        )
    widths = [18, 16, 10, 8, 50, 20]
    for index, width in enumerate(widths, start=1):
        note_sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(file_path)
    return file_path


def create_templates(output_dir: str | Path, fields: FieldsConfig) -> tuple[Path, Path]:
    directory = ensure_dir(Path(output_dir))
    csv_path = create_csv_template(directory / "sales_data_template.csv")
    xlsx_path = create_excel_template(directory / "sales_data_template.xlsx", fields)
    return csv_path, xlsx_path
