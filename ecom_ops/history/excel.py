"""历史数据 Excel 导出。"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ecom_ops.config import Settings
from ecom_ops.exceptions import HistoryError
from ecom_ops.history.query import (
    advice_effectiveness,
    metrics_for_export,
    overview,
    recent_runs,
    recommendations_for_export,
)
from ecom_ops.utils.dates import format_datetime
from ecom_ops.utils.files import ensure_dir


def _write_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    widths: list[int] | None = None,
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in rows:
        sheet.append(["" if value is None else value for value in row])
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
    for index, header in enumerate(headers, start=1):
        width = widths[index - 1] if widths and index <= len(widths) else min(
            36, max(12, len(header) * 2 + 4)
        )
        sheet.column_dimensions[get_column_letter(index)].width = width


def export_history_excel(
    settings: Settings,
    output_path: str | Path,
    days: int = 30,
) -> Path:
    db_path = settings.history.db_path
    if not db_path.exists():
        raise HistoryError(f"历史库不存在: {db_path}")

    info = overview(db_path)
    runs = recent_runs(db_path, limit=days)
    metrics = metrics_for_export(db_path, days=days)
    recommendations = recommendations_for_export(db_path, days=days)
    advice = advice_effectiveness(
        db_path,
        lookahead_days=settings.history.lookahead_days,
        limit=500,
    )
    anomalies = [item for item in metrics if item.get("anomaly_types")]

    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_sheet(
        workbook,
        "说明",
        ["项目", "内容"],
        [
            ["生成时间", format_datetime()],
            ["历史库", str(db_path)],
            ["导出天数", days],
            ["已记录天数", info["runs"]],
            ["日期范围", f"{info['first_date']} ~ {info['last_date']}"],
            ["涉及商品", info["products"]],
            ["累计异常", info["anomalies"]],
            ["累计建议", info["recommendations"]],
            ["AI 成功天数", info["ai_runs"]],
            ["数据说明", "历史复盘为观察性分析，不代表因果；请结合业务人工判断。"],
            ["隐私说明", "本文件只包含商品级指标与建议文本，不包含客户个人信息。"],
        ],
        widths=[18, 80],
    )

    _write_sheet(
        workbook,
        "每日概览",
        [
            "统计日期",
            "浏览量",
            "订单量",
            "销售额",
            "整体转化率",
            "浏览量变化率",
            "商品数",
            "异常数",
            "AI 状态",
            "数据说明",
            "报告路径",
        ],
        (
            [
                item["stat_date"],
                item["total_pv"],
                item["total_orders"],
                item["total_sales"],
                item["overall_cr"],
                item["overall_pv_change_rate"],
                item["product_count"],
                item["anomaly_count"],
                item["ai_status"],
                item["data_note"],
                item["report_path"],
            ]
            for item in runs
        ),
        widths=[12, 12, 10, 14, 12, 14, 10, 10, 16, 30, 50],
    )

    _write_sheet(
        workbook,
        "商品指标",
        [
            "统计日期",
            "商品ID",
            "商品名称",
            "店铺",
            "类目",
            "浏览量",
            "访客数",
            "订单量",
            "销售额",
            "转化率",
            "浏览量变化率",
            "基线浏览量",
            "基线订单",
            "异常类型",
            "严重度",
            "严重分",
            "低置信",
        ],
        (
            [
                item["stat_date"],
                item["product_id"],
                item["product_name"],
                item["shop"],
                item["category"],
                item["pv"],
                item["uv"],
                item["orders"],
                item["sales_amount"],
                item["cr"],
                item["pv_change_rate"],
                item["pv_baseline"],
                item["orders_baseline"],
                item["anomaly_types"],
                item["severity_level"],
                item["severity_score"],
                "是" if item["low_confidence"] else "否",
            ]
            for item in metrics
        ),
        widths=[12, 22, 36, 24, 16, 10, 10, 10, 12, 12, 14, 14, 12, 24, 10, 10, 8],
    )

    _write_sheet(
        workbook,
        "异常明细",
        [
            "统计日期",
            "商品ID",
            "商品名称",
            "店铺",
            "异常类型",
            "严重度",
            "严重分",
            "浏览量",
            "订单量",
            "转化率",
            "浏览量变化率",
        ],
        (
            [
                item["stat_date"],
                item["product_id"],
                item["product_name"],
                item["shop"],
                item["anomaly_types"],
                item["severity_level"],
                item["severity_score"],
                item["pv"],
                item["orders"],
                item["cr"],
                item["pv_change_rate"],
            ]
            for item in anomalies
        ),
        widths=[12, 22, 36, 24, 28, 10, 10, 10, 10, 12, 14],
    )

    _write_sheet(
        workbook,
        "建议清单",
        ["统计日期", "商品ID", "商品名称", "店铺", "来源", "建议", "理由", "风险", "需人工确认"],
        (
            [
                item["stat_date"],
                item["product_id"],
                item["product_name"],
                item["shop"],
                item["source"],
                item["action"],
                item["rationale"],
                item["risk"],
                "是" if item["needs_human"] else "否",
            ]
            for item in recommendations
        ),
        widths=[12, 22, 36, 24, 8, 45, 30, 30, 10],
    )

    _write_sheet(
        workbook,
        "建议追踪",
        [
            "建议日期",
            "商品ID",
            "商品名称",
            "来源",
            "建议",
            "基线订单",
            "后续日期",
            "后续订单",
            "订单变化",
            "转化率变化",
            "观察结果",
        ],
        (
            [
                item["stat_date"],
                item["product_id"],
                item["product_name"],
                item["source"],
                item["action"],
                item["base_orders"],
                item["followup_date"],
                item["followup_orders"],
                item["order_change"],
                item["cr_change"],
                item["classification"],
            ]
            for item in advice["details"]
        ),
        widths=[12, 22, 36, 8, 45, 12, 12, 12, 12, 14, 12],
    )

    path = Path(output_path)
    ensure_dir(path.parent)
    try:
        workbook.save(path)
    except Exception as exc:
        raise HistoryError(f"Excel 导出失败: {path}\n{exc}") from exc
    return path
